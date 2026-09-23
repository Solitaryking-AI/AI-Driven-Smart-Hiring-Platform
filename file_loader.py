"""
Resume File Loader
-------------------
Extracts raw text from PDF and DOCX resumes.
"""

import pymupdf  # PyMuPDF (the `fitz` alias is deprecated)
import docx


def extract_text_from_pdf(pdf_path: str) -> str:
    """Extract raw text from a PDF resume."""
    text = ""
    with pymupdf.open(pdf_path) as doc:
        for page in doc:
            text += page.get_text()
    return text


def extract_text_from_docx(docx_path: str) -> str:
    """Extract raw text from a DOCX resume, including table cells
    (many resumes put contact info / skills in tables)."""
    doc = docx.Document(docx_path)
    parts = [para.text for para in doc.paragraphs]

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    return "\n".join(parts)


def extract_text_from_txt(txt_path: str) -> str:
    """Extract raw text from a plain-text (.txt) resume.

    Tries UTF-8 first (with BOM handling) since that covers the vast majority
    of resumes; falls back to Latin-1 (which accepts any byte sequence) so an
    unusual encoding never raises instead of just parsing imperfectly.
    """
    try:
        with open(txt_path, "r", encoding="utf-8-sig") as f:
            return f.read()
    except UnicodeDecodeError:
        with open(txt_path, "r", encoding="latin-1") as f:
            return f.read()


def extract_text(file_path: str) -> str:
    """Dispatch to the right extractor based on file extension."""
    lower = file_path.lower()
    if lower.endswith(".pdf"):
        return extract_text_from_pdf(file_path)
    elif lower.endswith(".docx"):
        return extract_text_from_docx(file_path)
    elif lower.endswith(".txt"):
        return extract_text_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_path}. Only .pdf, .docx, and .txt are supported.")
